import streamlit as st
import pandas as pd
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from core.audit import comparer_balances

st.set_page_config(page_title="Audit - Intangibilité du bilan", layout="wide")
st.title("📊 Audit - Intangibilité du bilan")

file_n1 = st.file_uploader("Balance N-1", type=["csv", "xlsx"])
file_n = st.file_uploader("Balance N", type=["csv", "xlsx"])


def load_file(file):
    if file.name.endswith(".csv"):
        return pd.read_csv(file)
    return pd.read_excel(file)


def style_ecarts(val):
    if pd.isna(val):
        return ""
    if val != 0:
        return "color: red; font-weight: bold;"
    return ""


def build_excel_report(df: pd.DataFrame) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "CONTROLE"

    # écrire en-têtes
    ws.append(list(df.columns))

    # écrire données
    for row in df.itertuples(index=False, name=None):
        ws.append(list(row))

    # largeur colonnes
    widths = {
        "A": 15, "B": 35, "C": 15, "D": 35, "E": 15, "F": 15, "G": 30
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    # format nombres
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=6):
        for cell in row:
            cell.number_format = '#,##0.00'

    # rouge sur écarts non nuls
    red_fill = PatternFill(fill_type="solid", fgColor="FFC7CE")
    red_font = Font(color="9C0006", bold=True)

    for row in range(2, ws.max_row + 1):
        cell = ws[f"F{row}"]  # ECART
        if cell.value not in (None, 0):
            cell.fill = red_fill
            cell.font = red_font

    # créer un vrai tableau Excel
    table_ref = f"A1:G{ws.max_row}"
    table = Table(displayName="TableauControle", ref=table_ref)
    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    table.tableStyleInfo = style
    ws.add_table(table)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


if file_n1 is None or file_n is None:
    st.info("Veuillez importer les deux fichiers.")
    st.stop()

df_n1 = load_file(file_n1)
df_n = load_file(file_n)

result = comparer_balances(df_n1, df_n)

st.subheader("Résultat")
st.dataframe(result.style.map(style_ecarts, subset=["ECART"]), use_container_width=True)

anomalies = result[result["MESSAGE"] != "OK"].copy()

st.subheader("Anomalies")
st.dataframe(anomalies.style.map(style_ecarts, subset=["ECART"]), use_container_width=True)

st.metric("Nombre d'anomalies", len(anomalies))

excel_file = build_excel_report(result)

st.download_button(
    label="📥 Télécharger le rapport Excel",
    data=excel_file,
    file_name="rapport_intangibilite.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
)
